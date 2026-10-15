"""
backend/importer/wechat.py —— 微信支付账单 (.xlsx) 解析

真实文件结构（已实测确认）：
    第  0 行  '微信支付账单明细'
    第  1 行  '微信昵称：[Funny Ryan]'
    第  2 行  '起始时间：[...] 终止时间：[...]'
    第  3~9 行 导出类型 / 导出时间 / 共N笔记录 / 收入统计 / 支出统计 / 中性交易统计
    第 10~15 行 注意事项
    第 16 行  '----------------------微信支付账单明细列表--------------------'
    第 17 行  ★ 真正的表头
    第 18 行起 数据

⚠️ 表头行号不固定（取决于元信息有几行），必须【搜索】而不是硬编码。
"""
from __future__ import annotations

from pathlib import Path
from typing import Iterator

import openpyxl

from .base import (
    Row,
    clean,
    norm_amount,
    norm_direction,
    norm_time,
    should_drop,
)

# 表头关键词 → 统一字段名
HEADER_MAP = {
    "交易时间": "tx_time",
    "交易类型": "tx_type",
    "交易对方": "counterparty",
    "商品": "item",
    "收/支": "direction",
    "金额(元)": "amount",
    "金额（元）": "amount",
    "支付方式": "pay_method",
    "当前状态": "status",
    "交易单号": "tx_no",
    "商户单号": "merchant_no",
    "备注": "remark",
}


def find_header_row(ws, max_scan: int = 40) -> int:
    """搜索真正的表头行（第一列 == '交易时间'）。

    不能硬编码 17 —— 不同导出时间/版本的元信息行数会变。
    """
    for i, row in enumerate(ws.iter_rows(max_row=max_scan, values_only=True)):
        if row and clean(row[0]) == "交易时间":
            return i
    raise ValueError("未找到表头行：前 40 行内没有 '交易时间'")


def parse(path: Path) -> Iterator[Row]:
    """解析微信 xlsx，产出归一化 Row。"""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        hdr_idx = find_header_row(ws)

        header = [clean(c) for c in next(
            ws.iter_rows(min_row=hdr_idx + 1, max_row=hdr_idx + 1, values_only=True)
        )]
        col = {}
        for i, name in enumerate(header):
            if name in HEADER_MAP:
                col[HEADER_MAP[name]] = i

        missing = {"tx_time", "direction", "amount"} - set(col)
        if missing:
            raise ValueError(f"微信账单缺少必需列: {missing}")

        for raw_row in ws.iter_rows(min_row=hdr_idx + 2, values_only=True):
            if not raw_row or raw_row[0] is None:
                continue

            def get(key: str):
                idx = col.get(key)
                return raw_row[idx] if idx is not None and idx < len(raw_row) else None

            status = clean(get("status"))
            if should_drop(status):
                continue

            tx_time = norm_time(get("tx_time"))
            amount = norm_amount(get("amount"))
            if tx_time is None or amount is None:
                continue

            raw_dict = {
                header[i]: (str(raw_row[i]) if i < len(raw_row) and raw_row[i] is not None else "")
                for i in range(len(header))
            }

            yield Row(
                platform_id="wechat",
                source="微信支付",
                tx_time=tx_time,
                direction=norm_direction(get("direction")),
                amount=amount,
                pay_method=clean(get("pay_method")),
                counterparty=clean(get("counterparty")),
                item=clean(get("item")),
                tx_type=clean(get("tx_type")),
                status=status,
                tx_no=clean(get("tx_no")),
                merchant_no=clean(get("merchant_no")),
                remark=clean(get("remark")),
                raw=raw_dict,
            )
    finally:
        wb.close()
